import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont

wb = openpyxl.Workbook()

# ============================================================
# Sheet 1: 运镜测试表 (Camera Movement Test Matrix)
# ============================================================
ws = wb.active
ws.title = "运镜测试表"

header_font = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
normal_font = Font(name="Microsoft YaHei", size=10)
kw_font = Font(name="Microsoft YaHei", bold=True, size=10, color="C55A11")
kw_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")
center_align = Alignment(wrap_text=True, vertical="center", horizontal="center")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

bold_orange = InlineFont(b=True, color="C55A11")
normal_inline = InlineFont()

headers = [
    "序号", "运镜大类", "运镜子类", "中文描述",
    "室内完整Prompt", "室外完整Prompt",
    "英文运镜关键词", "英文运镜描述",
    "室内英文Prompt", "室外英文Prompt",
    "时长",
    "室内测试结果\n(1-5分)", "室外测试结果\n(1-5分)",
    "是否可识别", "备注"
]

en_header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = en_header_fill if col in (7, 8, 9, 10) else header_fill
    cell.alignment = center_align
    cell.border = thin_border

col_widths = [5, 12, 16, 18, 48, 48, 26, 52, 56, 56, 6, 12, 12, 10, 18]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

base_subject = "一个女人手里拿着一瓶香水对着镜头展示"
en_base = "a woman holding a perfume bottle presenting to camera"

# (大类, 子类, 中文描述, 室内prompt, 室外prompt, 英文关键词)
test_cases = [
    ("固定镜头", "纯固定", "镜头完全不动，只有人物动作",
     f"固定镜头，中近景，{base_subject}，室内明亮的化妆台前，柔和的环形灯光，浅景深背景虚化，镜头完全静止不动，5秒",
     f"固定镜头，中近景，{base_subject}，户外花园中，自然光下午阳光，绿植背景虚化，镜头完全静止不动，5秒",
     "static shot, locked-off"),
    ("固定镜头", "微固定(呼吸感)", "极轻微抖动增加真实感",
     f"固定镜头带轻微呼吸感，中近景，{base_subject}，室内白色简约背景，柔光灯，几乎不可察觉的微小晃动，5秒",
     f"固定镜头带轻微呼吸感，中近景，{base_subject}，户外咖啡厅露台，午后暖阳，几乎不可察觉的微小晃动，5秒",
     "static with subtle breathing"),
    ("推镜头", "缓慢推进", "镜头缓慢向人物靠近",
     f"缓慢推镜头，从中景推至中近景，{base_subject}，室内奢华梳妆台前，暖色台灯光，镜头缓缓向人物靠近，5秒",
     f"缓慢推镜头，从中景推至中近景，{base_subject}，户外日落海边，金色夕阳侧光，镜头缓缓向人物靠近，5秒",
     "slow dolly in"),
    ("推镜头", "快速推近", "镜头快速向人物靠近",
     f"快速推镜头，从中景快速推至近景，{base_subject}，室内时尚摄影棚，闪光灯效果，镜头迅速向人物推进，5秒",
     f"快速推镜头，从中景快速推至近景，{base_subject}，户外都市街头，霓虹灯光，镜头迅速向人物推进，5秒",
     "fast push in"),
    ("推镜头", "推至特写", "从中景推到香水瓶特写",
     f"推镜头至特写，{base_subject}，镜头从人物中景缓慢推进到香水瓶的极近特写，室内大理石桌面，柔和侧光，5秒",
     f"推镜头至特写，{base_subject}，镜头从人物中景缓慢推进到香水瓶的极近特写，户外阳光下草地，自然光，5秒",
     "dolly in to close-up"),
    ("拉镜头", "缓慢拉远", "镜头缓慢远离人物",
     f"缓慢拉镜头，从近景拉至中景，{base_subject}，室内现代客厅沙发旁，落地窗自然光，镜头缓慢后退，5秒",
     f"缓慢拉镜头，从近景拉至中景，{base_subject}，户外樱花树下，柔和散射光，镜头缓慢后退，5秒",
     "slow dolly out"),
    ("拉镜头", "拉远后固定", "拉到一定距离后停住",
     f"拉镜头后固定，{base_subject}，镜头从近景后退至中景后停住固定，室内书房，台灯暖光，5秒",
     f"拉镜头后固定，{base_subject}，镜头从近景后退至中景后停住固定，户外公园长椅旁，午后阳光，5秒",
     "dolly out then hold"),
    ("拉镜头", "快速拉远", "镜头快速远离主体",
     f"快速拉镜头，{base_subject}，镜头从近景快速拉至远景揭示整个房间，室内豪华酒店套房，水晶灯光，5秒",
     f"快速拉镜头，{base_subject}，镜头从近景快速拉至远景揭示整个场景，户外海边悬崖，日落逆光，5秒",
     "fast pull out"),
    ("横移镜头", "向左平移", "镜头匀速向左水平移动",
     f"横移镜头向左，中景，{base_subject}，摄影机从右向左匀速水平移动，室内长廊背景，均匀柔光，5秒",
     f"横移镜头向左，中景，{base_subject}，摄影机从右向左匀速水平移动，户外街道背景，自然光，5秒",
     "truck left, tracking left"),
    ("横移镜头", "向右平移", "镜头匀速向右水平移动",
     f"横移镜头向右，中景，{base_subject}，摄影机从左向右匀速水平移动，室内展示柜前，射灯照明，5秒",
     f"横移镜头向右，中景，{base_subject}，摄影机从左向右匀速水平移动，户外花墙背景，散射阳光，5秒",
     "truck right, tracking right"),
    ("横移镜头", "斜向移动", "镜头沿对角线方向移动",
     f"斜向横移镜头，中景，{base_subject}，镜头沿对角线从左下向右上移动，室内阶梯旁，侧光，5秒",
     f"斜向横移镜头，中景，{base_subject}，镜头沿对角线从左下向右上移动，户外斜坡草地，阳光，5秒",
     "diagonal tracking"),
    ("环绕镜头", "顺时针半环绕", "镜头绕主体顺时针转180°",
     f"环绕镜头，中景，{base_subject}，镜头从正面顺时针绕人物旋转180度到背面，室内纯白背景摄影棚，均匀灯光，5秒",
     f"环绕镜头，中景，{base_subject}，镜头从正面顺时针绕人物旋转180度到背面，户外海边沙滩，日落暖光，5秒",
     "180-degree clockwise orbit"),
    ("环绕镜头", "逆时针半环绕", "镜头绕主体逆时针转180°",
     f"环绕镜头，中景，{base_subject}，镜头从正面逆时针绕人物旋转180度，室内复古风房间，窗帘透光，5秒",
     f"环绕镜头，中景，{base_subject}，镜头从正面逆时针绕人物旋转180度，户外古建筑前，柔和天光，5秒",
     "180-degree counterclockwise orbit"),
    ("环绕镜头", "全环绕360°", "镜头绕主体旋转一整圈",
     f"360度环绕镜头，中景，{base_subject}，镜头绕人物匀速旋转一整圈，室内环形灯光摄影棚，5秒",
     f"360度环绕镜头，中景，{base_subject}，镜头绕人物匀速旋转一整圈，户外草原，蓝天白云，5秒",
     "360-degree orbit"),
    ("升降镜头", "升镜头", "镜头从低位上升到高位",
     f"升镜头，{base_subject}，镜头从脚部低位缓慢上升至面部平视位，室内大厅，水晶吊灯，5秒",
     f"升镜头，{base_subject}，镜头从脚部低位缓慢上升至面部平视位，户外台阶前，正午顶光，5秒",
     "crane up, boom up, rising shot"),
    ("升降镜头", "降镜头", "镜头从高位下降到低位",
     f"降镜头，{base_subject}，镜头从高处俯视缓慢下降至面部平视位，室内楼梯间，自然采光，5秒",
     f"降镜头，{base_subject}，镜头从高处俯视缓慢下降至面部平视位，户外阳台，城市天际线背景，5秒",
     "crane down, boom down, descending shot"),
    ("升降镜头", "升后俯瞰", "升到最高后俯视主体",
     f"升镜头至俯瞰，{base_subject}，镜头从平视缓慢上升至鸟瞰视角俯视，室内圆形大厅，顶光，5秒",
     f"升镜头至俯瞰，{base_subject}，镜头从平视缓慢上升至鸟瞰视角俯视，户外广场，阳光投射阴影，5秒",
     "crane up to bird's eye"),
    ("摇镜头", "水平右摇", "机位固定，镜头从左摇到右",
     f"水平右摇，{base_subject}，摄像机固定不动，镜头从左侧慢慢摇到右侧揭示人物，室内化妆间，镜面反射光，5秒",
     f"水平右摇，{base_subject}，摄像机固定不动，镜头从左侧慢慢摇到右侧揭示人物，户外湖边，水面反光，5秒",
     "pan right"),
    ("摇镜头", "水平左摇", "机位固定，镜头从右摇到左",
     f"水平左摇，{base_subject}，摄像机固定不动，镜头从右侧慢慢摇到左侧揭示人物，室内展厅，聚光灯，5秒",
     f"水平左摇，{base_subject}，摄像机固定不动，镜头从右侧慢慢摇到左侧揭示人物，户外街角，路灯暖光，5秒",
     "pan left"),
    ("摇镜头", "垂直上摇", "机位固定，镜头从下向上摇",
     f"垂直上摇，{base_subject}，摄像机固定不动，镜头从香水瓶向上摇至人物面部，室内简约背景，顶部柔光，5秒",
     f"垂直上摇，{base_subject}，摄像机固定不动，镜头从香水瓶向上摇至人物面部，户外瀑布前，自然光，5秒",
     "tilt up"),
    ("摇镜头", "垂直下摇", "机位固定，镜头从上向下摇",
     f"垂直下摇，{base_subject}，摄像机固定不动，镜头从面部向下摇至香水瓶，室内灰色调背景，环形灯，5秒",
     f"垂直下摇，{base_subject}，摄像机固定不动，镜头从面部向下摇至香水瓶，户外花丛中，晨光，5秒",
     "tilt down"),
    ("甩镜头", "水平甩", "镜头急速水平摇转",
     f"水平甩镜头，{base_subject}，镜头从一侧急速甩向人物方向，产生运动模糊，室内时尚背景，闪灯，5秒",
     f"水平甩镜头，{base_subject}，镜头从一侧急速甩向人物方向，产生运动模糊，户外都市天台，夜景灯光，5秒",
     "whip pan, swish pan"),
    ("甩镜头", "垂直甩", "镜头急速垂直摇转",
     f"垂直甩镜头，{base_subject}，镜头从上方急速甩向下方人物位置，产生运动模糊，室内挑高空间，吊灯，5秒",
     f"垂直甩镜头，{base_subject}，镜头从上方急速甩向下方人物位置，产生运动模糊，户外高楼间，天空背景，5秒",
     "vertical whip tilt"),
    ("晃动镜头", "轻微晃动(手持感)", "轻微抖动增加真实感",
     f"手持轻微晃动，中近景，{base_subject}，带有呼吸感的自然轻微晃动，室内日式和风房间，柔光，5秒",
     f"手持轻微晃动，中近景，{base_subject}，带有呼吸感的自然轻微晃动，户外竹林小路，斑驳树影，5秒",
     "subtle handheld"),
    ("晃动镜头", "明显晃动(纪录片感)", "明显抖动纪实效果",
     f"明显手持晃动，中景，{base_subject}，摄像机有明显的手持抖动感，室内工作室后台，荧光灯，5秒",
     f"明显手持晃动，中景，{base_subject}，摄像机有明显的手持抖动感，户外市集人群中，自然光，5秒",
     "shaky handheld"),
    ("变焦镜头", "缓慢变焦推", "焦距缓慢拉近",
     f"缓慢变焦推进，中景至近景，{base_subject}，通过变焦缓慢放大画面，室内纯色背景，柔和灯光，5秒",
     f"缓慢变焦推进，中景至近景，{base_subject}，通过变焦缓慢放大画面，户外山顶观景台，远景虚化，5秒",
     "slow zoom in"),
    ("变焦镜头", "快速变焦推(甩焦)", "焦距瞬间拉近",
     f"快速变焦推，{base_subject}，镜头瞬间从中景变焦至面部特写，室内派对场景，彩色灯光，5秒",
     f"快速变焦推，{base_subject}，镜头瞬间从中景变焦至面部特写，户外音乐节，舞台灯光，5秒",
     "snap zoom in"),
    ("变焦镜头", "变焦拉(缩小)", "焦距缓慢拉远",
     f"变焦拉，{base_subject}，通过变焦缓慢缩小画面从近景到中景，室内落地窗旁，侧逆光，5秒",
     f"变焦拉，{base_subject}，通过变焦缓慢缩小画面从近景到中景，户外桥上，河面反光，5秒",
     "zoom out"),
    ("跟随镜头", "正面后退跟", "镜头在前方后退拍正面",
     f"正面跟随镜头，中景，{base_subject}并缓慢向前走来，镜头在正前方后退拍摄，室内长走廊，两侧壁灯，5秒",
     f"正面跟随镜头，中景，{base_subject}并缓慢向前走来，镜头在正前方后退拍摄，户外林荫道，斑驳阳光，5秒",
     "camera moves backward facing subject"),
    ("跟随镜头", "侧面跟随", "镜头在侧面平行跟随",
     f"侧面跟随镜头，中景，{base_subject}并向右行走，镜头在侧面平行移动，室内商场走廊，明亮灯光，5秒",
     f"侧面跟随镜头，中景，{base_subject}并向右行走，镜头在侧面平行移动，户外海滨步道，海风吹发，5秒",
     "side tracking shot"),
    ("滚转镜头", "顺时针滚转", "画面沿光轴顺时针旋转",
     f"滚转镜头，中近景，{base_subject}，画面缓慢顺时针旋转，室内暗色调背景，聚光灯，梦幻效果，5秒",
     f"滚转镜头，中近景，{base_subject}，画面缓慢顺时针旋转，户外天空背景，飘浮云朵，梦幻效果，5秒",
     "clockwise roll"),
    ("滚转镜头", "逆时针滚转", "画面沿光轴逆时针旋转",
     f"滚转镜头，中近景，{base_subject}，画面缓慢逆时针旋转，室内镜面反射背景，霓虹灯，5秒",
     f"滚转镜头，中近景，{base_subject}，画面缓慢逆时针旋转，户外星空背景，月光，5秒",
     "counterclockwise roll"),
    ("复合运镜", "推+升", "靠近主体同时镜头上升",
     f"推镜头同时升镜头，{base_subject}，镜头从低位中景缓慢靠近并上升至面部平视，室内剧场舞台，追光灯，5秒",
     f"推镜头同时升镜头，{base_subject}，镜头从低位中景缓慢靠近并上升至面部平视，户外古堡台阶前，日落光，5秒",
     "dolly in while craning up"),
    ("复合运镜", "拉+降", "远离主体同时镜头下降",
     f"拉镜头同时降镜头，{base_subject}，镜头从面部平视缓慢后退并下降，室内酒店大堂，水晶灯，5秒",
     f"拉镜头同时降镜头，{base_subject}，镜头从面部平视缓慢后退并下降，户外教堂前广场，阴天柔光，5秒",
     "dolly out while descending"),
    ("复合运镜", "横移+推", "先侧移再转为推进",
     f"先横移后推进，中景，{base_subject}，镜头先向右横移再转为正面推进至近景，室内画廊，射灯，5秒",
     f"先横移后推进，中景，{base_subject}，镜头先向右横移再转为正面推进至近景，户外壁画墙前，自然光，5秒",
     "lateral tracking to dolly in"),
    ("复合运镜", "环绕+升", "环绕主体同时升高",
     f"螺旋上升环绕，中景，{base_subject}，镜头绕人物旋转同时逐渐升高至俯视，室内圆形大厅，环形灯，5秒",
     f"螺旋上升环绕，中景，{base_subject}，镜头绕人物旋转同时逐渐升高至俯视，户外圆形花坛中央，阳光，5秒",
     "spiral ascending orbit"),
    ("复合运镜", "希区柯克变焦", "推镜头+反向变焦产生眩晕",
     f"希区柯克变焦效果，中景，{base_subject}，镜头向前推进同时焦距拉远，背景产生拉伸变形，室内走廊，荧光灯，5秒",
     f"希区柯克变焦效果，中景，{base_subject}，镜头向前推进同时焦距拉远，背景产生拉伸变形，户外隧道出口，逆光，5秒",
     "dolly zoom, vertigo effect"),
    ("横移镜头", "遮挡横移", "藏在柱子/门后制造悬疑惊喜",
     f"遮挡横移镜头，中景，{base_subject}，镜头从柱子后方水平移出揭示人物再移入柱子后方消失，室内大厅立柱旁，侧光，5秒",
     f"遮挡横移镜头，中景，{base_subject}，镜头从树干后方水平移出揭示人物再移入另一棵树后方，户外林荫道，斑驳阳光，5秒",
     "obstacle tracking"),
    ("环绕镜头", "缓慢弧形环绕", "缓慢弧线转换视角唯美治愈",
     f"缓慢弧形环绕，中景，{base_subject}，镜头以缓慢弧线绕人物移动转换视角，室内花艺工作室，柔和自然光，唯美风格，5秒",
     f"缓慢弧形环绕，中景，{base_subject}，镜头以缓慢弧线绕人物移动转换视角，户外薰衣草花田，金色夕阳，唯美治愈，5秒",
     "slow arc orbit"),
    ("摇镜头", "仰拍上摇", "低角度从脚拍到脸提升气场",
     f"仰拍上摇，{base_subject}，镜头从低角度脚部开始向上摇至面部，提升人物气场，室内大理石地面，顶光，5秒",
     f"仰拍上摇，{base_subject}，镜头从低角度脚部开始向上摇至面部，提升人物气场，户外台阶底部仰视，逆光剪影，5秒",
     "low-angle tilt up"),
    ("摇镜头", "俯拍下摇", "高角度从上拍到全身展示穿搭",
     f"俯拍下摇，{base_subject}，镜头从高角度面部开始向下摇至全身，展示完整造型，室内衣帽间，明亮灯光，5秒",
     f"俯拍下摇，{base_subject}，镜头从高角度面部开始向下摇至全身，展示完整造型，户外天台，城市天际线背景，5秒",
     "high-angle tilt down"),
    ("晃动镜头", "手持纪实", "微抖有生活感和纪录片感",
     f"手持纪实风格，中近景，{base_subject}，带有纪录片式的微微抖动，室内厨房窗边，自然光，生活气息，5秒",
     f"手持纪实风格，中近景，{base_subject}，带有纪录片式的微微抖动，户外早市街边，晨光，烟火气，5秒",
     "documentary handheld"),
    ("变焦镜头", "微距变焦", "变焦至极致细节有科幻微观感",
     f"微距变焦，{base_subject}，镜头变焦推进至香水瓶表面纹理的微距特写，室内暗色背景，聚光灯，科幻感，5秒",
     f"微距变焦，{base_subject}，镜头变焦推进至香水瓶表面水珠的微距特写，户外雨后，水滴折射阳光，5秒",
     "macro zoom"),
    ("变焦镜头", "急速冲撞变焦", "瞬间特写表现震惊冲击",
     f"急速冲撞变焦，中景至特写，{base_subject}，镜头瞬间冲向人物面部产生强烈冲击感，室内深色背景，戏剧性灯光，5秒",
     f"急速冲撞变焦，中景至特写，{base_subject}，镜头瞬间冲向人物面部产生强烈冲击感，户外空旷地，强烈日光，5秒",
     "crash zoom"),
    ("变焦镜头", "平滑光学推进", "温柔拉近突出面部情绪",
     f"平滑光学推进，中景至近景，{base_subject}，镜头通过光学变焦温柔平滑地拉近人物面部，室内暖色调客厅，落地窗柔光，5秒",
     f"平滑光学推进，中景至近景，{base_subject}，镜头通过光学变焦温柔平滑地拉近人物面部，户外樱花树下，花瓣飘落，5秒",
     "smooth optical zoom in"),
    ("变焦镜头", "平滑光学拉远", "慢慢拉宽显环境空旷",
     f"平滑光学拉远，近景至中景，{base_subject}，镜头通过光学变焦缓慢拉远显示更多环境，室内空旷画廊，回音感，5秒",
     f"平滑光学拉远，近景至中景，{base_subject}，镜头通过光学变焦缓慢拉远显示更多环境，户外沙漠公路，广袤天际，5秒",
     "smooth optical zoom out"),
    ("滚转镜头", "荷兰倾斜角", "画面歪斜表现紧张不安",
     f"荷兰倾斜角，中近景，{base_subject}，画面保持倾斜角度拍摄产生不安感，室内走廊尽头，荧光灯闪烁，悬疑氛围，5秒",
     f"荷兰倾斜角，中近景，{base_subject}，画面保持倾斜角度拍摄产生不安感，户外废弃工厂前，阴天灰色调，5秒",
     "dutch angle"),
    ("跟随镜头", "第一人称行走", "沉浸式第一人称视角超强代入",
     f"第一人称视角，{base_subject}并向前走动，镜头模拟人物第一视角看到的画面，室内商场走廊，明亮灯光，沉浸感，5秒",
     f"第一人称视角，{base_subject}并向前走动，镜头模拟人物第一视角看到的画面，户外森林小路，阳光透过树叶，5秒",
     "first-person walking"),
    ("航拍运镜", "俯冲", "从高空快速下降至主体",
     f"航拍俯冲，{base_subject}站在中央，无人机从高空快速下降俯冲至人物近景，室内超大中庭，玻璃穹顶采光，5秒",
     f"航拍俯冲，{base_subject}站在草地中央，无人机从高空快速下降俯冲至人物近景，户外山顶草坪，蓝天白云，5秒",
     "drone dive"),
    ("航拍运镜", "拉升", "从低位快速上升至鸟瞰",
     f"航拍拉升，{base_subject}，镜头从人物近景快速上升至鸟瞰全景，室内体育馆中央，顶光照射，5秒",
     f"航拍拉升，{base_subject}站在海边，镜头从人物近景快速上升至鸟瞰海岸线全景，户外碧海蓝天，5秒",
     "drone ascend, aerial pull up"),
    ("航拍运镜", "无人机高空飞跃", "上帝视角飞越开阔治愈",
     f"无人机高空飞跃，{base_subject}站在中央，镜头从高空匀速飞越人物上方，室内大型展厅，鸟瞰全景，5秒",
     f"无人机高空飞跃，{base_subject}站在田野中，镜头从高空匀速飞越人物上方，户外金色麦田，开阔天空，5秒",
     "aerial flyover"),
    ("航拍运镜", "史诗无人机揭示", "从遮挡物后露出惊艳开场",
     f"史诗无人机揭示，{base_subject}站在远处，镜头从建筑物后方升起揭示人物和整个大厅，室内穹顶教堂，圣光，5秒",
     f"史诗无人机揭示，{base_subject}站在山顶，镜头从山脊后方升起揭示人物和壮阔山景，户外日出金光，5秒",
     "epic drone reveal"),
    ("航拍运镜", "上帝俯拍", "90度正俯视构图高级",
     f"上帝俯拍视角，{base_subject}仰头看镜头，90度正俯视角度，室内白色地面，几何阴影，极简构图，5秒",
     f"上帝俯拍视角，{base_subject}仰头看镜头，90度正俯视角度，户外沙滩上，海浪拍岸，航拍构图，5秒",
     "top-down, bird's eye view"),
    ("航拍运镜", "FPV穿越机", "高速穿越飞行刺激炸裂",
     f"FPV穿越机视角，{base_subject}站在尽头，镜头高速穿越柱子和障碍物飞向人物，室内仓库空间，工业灯光，5秒",
     f"FPV穿越机视角，{base_subject}站在桥头，镜头高速穿越树丛和栏杆飞向人物，户外森林吊桥，自然光，5秒",
     "FPV drone shot"),
    ("特殊视角", "过肩镜头", "从肩后拍摄增强代入感故事感",
     f"过肩镜头，{base_subject}，从另一人肩膀后方拍摄主体人物，室内办公室，柔和灯光，对话氛围，5秒",
     f"过肩镜头，{base_subject}，从另一人肩膀后方拍摄主体人物，户外咖啡厅露台，午后阳光，对话氛围，5秒",
     "over-the-shoulder (OTS)"),
    ("特殊视角", "鱼眼/猫眼视角", "画面扭曲制造压迫悬疑感",
     f"鱼眼视角，中近景，{base_subject}，鱼眼镜头产生画面边缘扭曲效果，室内门口猫眼视角，暗色调，悬疑感，5秒",
     f"鱼眼视角，中近景，{base_subject}，鱼眼镜头产生画面边缘扭曲效果，户外巷道尽头，阴影浓重，不安氛围，5秒",
     "fisheye, peephole"),
    ("焦点控制", "虚焦转清晰", "从模糊渐清晰营造回忆心动感",
     f"虚焦转清晰，{base_subject}，画面从完全虚焦逐渐对焦至人物面部清晰，室内梦幻背景，光斑散景，回忆感，5秒",
     f"虚焦转清晰，{base_subject}，画面从完全虚焦逐渐对焦至人物面部清晰，户外花田，阳光光斑，心动氛围，5秒",
     "rack focus in"),
    ("焦点控制", "焦点切换", "前后景焦点转移突出重点",
     f"焦点切换，{base_subject}，焦点从前景花瓶切换到背景人物面部，室内桌面摆设，浅景深，5秒",
     f"焦点切换，{base_subject}，焦点从前景树叶切换到背景人物面部，户外花园栅栏旁，自然光，5秒",
     "rack focus shift"),
]

# (en_movement_desc, shot_size, indoor_scene, outdoor_scene)
en_details = [
    ("Static shot, locked-off camera, completely still frame", "medium close-up",
     "bright vanity table, soft ring light, shallow depth of field",
     "garden, natural afternoon sunlight, blurred green background"),
    ("Static shot with subtle breathing, barely perceptible micro-shake", "medium close-up",
     "minimalist white background, soft fill light",
     "cafe terrace, warm afternoon sunlight"),
    ("Slow dolly in, camera steadily moves closer from medium shot to medium close-up", "medium shot to medium close-up",
     "luxurious vanity table, warm table lamp light",
     "sunset beach, golden side light"),
    ("Fast push in, camera rapidly advances from medium shot to close-up", "medium shot to close-up",
     "fashion studio, strobe light effect",
     "urban street, neon lights"),
    ("Dolly in to extreme close-up, camera slowly pushes to product detail", "medium shot to extreme close-up",
     "marble tabletop, soft side light",
     "sunlit grass field, natural light"),
    ("Slow dolly out, camera gently pulls back from close-up to medium shot", "close-up to medium shot",
     "modern living room, floor-to-ceiling window natural light",
     "under cherry blossom tree, soft diffused light"),
    ("Dolly out then hold, camera pulls back to medium shot then locks off", "close-up to medium shot",
     "study room, warm desk lamp",
     "park bench, afternoon sunlight"),
    ("Fast pull out, camera rapidly retreats from close-up to wide shot revealing scene", "close-up to wide shot",
     "luxury hotel suite, crystal chandelier",
     "seaside cliff, sunset backlight"),
    ("Truck left, camera moves horizontally from right to left at steady pace", "medium shot",
     "long corridor background, even soft light",
     "street background, natural light"),
    ("Truck right, camera moves horizontally from left to right at steady pace", "medium shot",
     "display cabinet, spotlight illumination",
     "flower wall background, scattered sunlight"),
    ("Diagonal tracking, camera moves along diagonal from lower-left to upper-right", "medium shot",
     "staircase area, side light",
     "sloped grassland, sunlight"),
    ("180-degree clockwise orbit, camera rotates around subject from front to back", "medium shot",
     "pure white studio, even lighting",
     "beach, warm sunset light"),
    ("180-degree counterclockwise orbit, camera revolves around subject", "medium shot",
     "vintage-style room, curtain-filtered light",
     "ancient architecture, soft ambient light"),
    ("360-degree orbit, camera smoothly revolves full circle around subject", "medium shot",
     "ring-lit studio",
     "grassland, blue sky with white clouds"),
    ("Crane up / boom up, camera rises from low position to eye level", "low angle to eye level",
     "grand hall, crystal chandelier",
     "front of steps, noon overhead light"),
    ("Crane down / boom down, camera descends from overhead to eye level", "high to eye level",
     "stairwell, natural daylight",
     "balcony, city skyline background"),
    ("Crane up to bird's eye view, camera rises to overhead looking down", "eye level to top-down",
     "circular hall, overhead light",
     "plaza, sunlight casting shadows"),
    ("Pan right, camera stays fixed rotating horizontally left to right to reveal subject", "medium shot",
     "dressing room, mirror reflection light",
     "lakeside, water surface glare"),
    ("Pan left, camera stays fixed rotating horizontally right to left to reveal subject", "medium shot",
     "exhibition hall, spotlight",
     "street corner, warm street lamp"),
    ("Tilt up, camera stays fixed tilting vertically from product up to face", "close-up to medium close-up",
     "minimal background, soft overhead light",
     "waterfall foreground, natural light"),
    ("Tilt down, camera stays fixed tilting vertically from face down to product", "medium close-up to close-up",
     "grey-tone background, ring light",
     "flower garden, morning light"),
    ("Horizontal whip pan, camera rapidly swings to subject with motion blur", "medium shot",
     "fashion backdrop, strobe light",
     "urban rooftop, night city lights"),
    ("Vertical whip tilt, camera rapidly swings downward to subject with motion blur", "medium shot",
     "high-ceiling space, pendant light",
     "between tall buildings, sky background"),
    ("Subtle handheld, gentle breathing-like micro-shake for natural feel", "medium close-up",
     "Japanese-style room, soft diffused light",
     "bamboo forest path, dappled shadows"),
    ("Shaky handheld, noticeable camera shake for documentary feel", "medium shot",
     "backstage workshop, fluorescent light",
     "crowded market, natural light"),
    ("Slow zoom in, focal length gradually increases magnifying frame", "medium shot to close-up",
     "solid color backdrop, soft lighting",
     "mountaintop viewpoint, distant haze"),
    ("Snap zoom in, instant focal length jump to face close-up", "medium shot to extreme close-up",
     "party scene, colorful lights",
     "music festival, stage lights"),
    ("Zoom out, focal length decreases revealing wider view", "close-up to medium shot",
     "beside floor-to-ceiling window, side backlight",
     "on a bridge, river reflection"),
    ("Tracking shot backward facing subject, camera retreats as subject walks forward", "medium shot",
     "long corridor, wall sconces both sides",
     "tree-lined avenue, dappled sunlight"),
    ("Side tracking shot, camera moves parallel alongside walking subject", "medium shot",
     "mall hallway, bright lighting",
     "coastal boardwalk, sea breeze blowing hair"),
    ("Clockwise roll, frame slowly rotates clockwise along lens axis", "medium close-up",
     "dark-toned backdrop, spotlight, dreamy effect",
     "sky background, floating clouds, dreamy effect"),
    ("Counterclockwise roll, frame slowly rotates counterclockwise along lens axis", "medium close-up",
     "mirror reflection backdrop, neon lights",
     "starry sky background, moonlight"),
    ("Dolly in while craning up, camera moves closer and rises simultaneously", "low medium shot to eye-level close-up",
     "theater stage, follow spot",
     "castle steps, sunset light"),
    ("Dolly out while descending, camera pulls back and lowers simultaneously", "eye-level close-up to low medium shot",
     "hotel lobby, crystal chandelier",
     "church plaza, overcast soft light"),
    ("Lateral tracking transitioning to dolly in, side movement shifts to forward push", "medium shot to close-up",
     "art gallery, spotlight",
     "mural wall, natural light"),
    ("Spiral ascending orbit, camera orbits subject while gradually rising to overhead", "medium shot to bird's eye",
     "circular hall, ring light",
     "circular flower bed center, sunlight"),
    ("Dolly zoom / vertigo effect, push forward while zooming out, background distorts", "medium shot",
     "corridor, fluorescent light",
     "tunnel exit, backlight"),
    ("Obstacle tracking, camera tracks behind pillars creating peek-a-boo reveal", "medium shot",
     "grand hall with pillars, side light",
     "tree-lined avenue, dappled sunlight"),
    ("Slow arc orbit, camera glides in gentle arc around subject, dreamy aesthetic", "medium shot",
     "floral workshop, soft natural light, aesthetic",
     "lavender field, golden sunset, healing vibe"),
    ("Low-angle tilt up, camera tilts from feet to face at low angle, empowering", "full shot to close-up low-angle",
     "marble floor, overhead spotlight",
     "bottom of steps, backlit silhouette"),
    ("High-angle tilt down, camera tilts from face to full body for outfit reveal", "close-up to full shot high-angle",
     "walk-in closet, bright even lighting",
     "rooftop terrace, city skyline backdrop"),
    ("Documentary handheld, subtle natural micro-shake for authentic lived-in feel", "medium close-up",
     "kitchen window, natural light, slice-of-life",
     "morning market street, dawn light, daily vibe"),
    ("Macro zoom, extreme close-up zoom revealing microscopic surface details", "medium close-up to extreme close-up macro",
     "dark background, spotlight, sci-fi feel",
     "after rain, water droplets refracting sunlight"),
    ("Crash zoom, instant snap to extreme close-up for shock impact", "medium shot to extreme close-up",
     "dark background, dramatic lighting",
     "open field, harsh direct sunlight"),
    ("Smooth optical zoom in, gentle focal length increase highlighting facial emotion", "medium shot to close-up",
     "warm-toned living room, window soft light",
     "under cherry blossoms, petals falling"),
    ("Smooth optical zoom out, gradual widening to reveal spacious environment", "close-up to medium shot",
     "spacious gallery, ambient echo",
     "desert highway, vast open horizon"),
    ("Dutch angle, tilted frame creating unease and visual tension", "medium close-up",
     "end of corridor, flickering fluorescent, suspenseful",
     "abandoned factory, overcast grey tones"),
    ("First-person walking POV, immersive subjective camera perspective", "first-person perspective",
     "mall corridor, bright lighting, immersive",
     "forest trail, sunlight filtering through leaves"),
    ("Drone dive, rapid aerial descent toward subject from high altitude", "aerial to close-up",
     "large indoor atrium, glass dome skylight",
     "hilltop grass, blue sky white clouds"),
    ("Drone ascend / aerial pull up, rising from subject to bird's eye overview", "close-up to aerial wide shot",
     "gymnasium center, overhead lighting",
     "seaside, revealing coastline panorama, blue ocean"),
    ("Aerial flyover, God's-eye view gliding steadily over the scene", "top-down aerial",
     "large exhibition hall, bird's eye panorama",
     "golden wheat field, expansive open sky"),
    ("Epic drone reveal, camera rises from behind obstacle to unveil grand vista", "reveal to wide shot",
     "domed cathedral interior, holy light beams",
     "mountain ridge, sunrise golden rays"),
    ("Top-down / bird's eye view, 90-degree overhead perspective straight down", "top-down",
     "white floor, geometric shadows, minimalist",
     "sandy beach, waves crashing, aerial composition"),
    ("FPV drone shot, high-speed first-person racing through obstacles", "first-person view continuous",
     "warehouse space, industrial lighting",
     "forest suspension bridge, natural light"),
    ("Over-the-shoulder shot (OTS), framing from behind another person's shoulder", "over-the-shoulder medium shot",
     "office setting, soft lighting, conversation",
     "outdoor cafe terrace, afternoon sunlight"),
    ("Fisheye / peephole perspective, distorted wide-angle, claustrophobic feel", "fisheye medium close-up",
     "doorway peephole angle, dark tones, suspense",
     "narrow alley, heavy shadows, uneasy atmosphere"),
    ("Rack focus in, transitioning from blurred bokeh to sharp focus on subject", "medium close-up rack focus",
     "dreamy background, circular bokeh orbs, nostalgic",
     "flower field, sun flare bokeh, heartfelt mood"),
    ("Rack focus shift, focus transfers from foreground element to background subject", "close-up rack focus",
     "tabletop props foreground, shallow depth of field",
     "garden fence foreground, natural light"),
]

# Reorder by platform recognition frequency (S → A → B → C → D)
_freq_order = [
    # S级: dolly in/out (#1)
    2, 3, 5, 6, 7,
    # S级: pan (#2)
    17, 18,
    # S级: zoom (#3)
    25, 26, 27,
    # S级: static (#4)
    0, 1,
    # A级: tilt (#5)
    19, 20,
    # A级: tracking (#6)
    28, 29,
    # A级: crane (#7)
    14, 15, 16,
    # A级: orbit (#8)
    11, 12, 13,
    # A级: handheld (#9)
    23, 24, 41,
    # A级: truck (#10)
    8, 9, 10,
    # A级: whip pan (#11)
    21, 22,
    # B级: push in (#12)
    4,
    # B级: dolly zoom (#15)
    36,
    # B级: crash zoom (#16)
    43,
    # B级: rack focus (#17)
    57,
    # B级: aerial (#18)
    48, 49, 50, 51,
    # B级: POV (#19)
    47,
    # B级: dutch angle (#20)
    46,
    # B级: compound movements
    32, 33, 34, 35,
    # C级: slow arc (#21)
    38,
    # C级: low-angle tilt (#23)
    39,
    # C级: high-angle tilt (#24)
    40,
    # C级: obstacle tracking (#25)
    37,
    # C级: macro zoom (#27)
    42,
    # C级: smooth optical zoom (#28)
    44, 45,
    # C级: OTS (#29)
    54,
    # C级: fisheye (#30)
    55,
    # D级: FPV (#31)
    53,
    # D级: roll (#32)
    30, 31,
    # D级: top-down (#33)
    52,
    # D级: rack focus in (#34)
    56,
]
test_cases = [test_cases[i] for i in _freq_order]
en_details = [en_details[i] for i in _freq_order]

TOTAL_COLS = 15

for idx, (cat, sub, desc, indoor, outdoor, en_kw) in enumerate(test_cases, 1):
    row = idx + 1
    mv_desc, ss, in_sc, out_sc = en_details[idx - 1]

    ws.cell(row=row, column=1, value=idx).font = normal_font
    ws.cell(row=row, column=2, value=cat).font = normal_font
    ws.cell(row=row, column=3, value=sub).font = normal_font
    ws.cell(row=row, column=4, value=desc).font = normal_font
    ws.cell(row=row, column=5, value=indoor).font = normal_font
    ws.cell(row=row, column=6, value=outdoor).font = normal_font

    c7 = ws.cell(row=row, column=7, value=en_kw)
    c7.font = kw_font
    c7.fill = kw_fill

    c8 = ws.cell(row=row, column=8, value=mv_desc)
    c8.font = kw_font
    c8.fill = kw_fill

    c9 = ws.cell(row=row, column=9)
    c9.value = CellRichText(
        TextBlock(bold_orange, mv_desc),
        TextBlock(normal_inline, f", {ss}, {en_base}, {in_sc}, 5 seconds"),
    )

    c10 = ws.cell(row=row, column=10)
    c10.value = CellRichText(
        TextBlock(bold_orange, mv_desc),
        TextBlock(normal_inline, f", {ss}, {en_base}, {out_sc}, 5 seconds"),
    )

    ws.cell(row=row, column=11, value="0-5s").font = normal_font
    for c in (12, 13, 14, 15):
        ws.cell(row=row, column=c, value="").font = normal_font

    for c in range(1, TOTAL_COLS + 1):
        ws.cell(row=row, column=c).alignment = wrap_align
        ws.cell(row=row, column=c).border = thin_border

# -- 景别缩写备注 --
note_font = Font(name="Microsoft YaHei", size=10)
note_title_font = Font(name="Microsoft YaHei", bold=True, size=11, color="2F5496")
note_header_font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
note_header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

note_start = len(test_cases) + 4  # 2 blank rows gap

ws.cell(row=note_start, column=1, value="📌 景别缩写对照表（Shot Size Abbreviations）").font = note_title_font
ws.merge_cells(start_row=note_start, start_column=1, end_row=note_start, end_column=5)

note_headers = ["缩写", "英文全称", "中文", "说明"]
for col, h in enumerate(note_headers, 1):
    cell = ws.cell(row=note_start + 1, column=col, value=h)
    cell.font = note_header_font
    cell.fill = note_header_fill
    cell.alignment = center_align
    cell.border = thin_border

shot_abbr = [
    ("ECU", "Extreme Close-Up", "大特写", "眼睛、嘴唇、产品纹理等极致细节"),
    ("CU",  "Close-Up", "特写/近景", "面部或产品占满画面"),
    ("MCU", "Medium Close-Up", "中近景", "胸部以上，最常用的口播景别"),
    ("MS",  "Medium Shot", "中景", "腰部以上，兼顾人物与环境"),
    ("FS",  "Full Shot", "全景", "完整人物从头到脚"),
    ("WS",  "Wide Shot", "远景", "人物较小，环境为主"),
    ("OTS", "Over-The-Shoulder", "过肩镜头", "从另一人肩后拍摄，对话常用"),
    ("POV", "Point of View", "主观视角", "模拟角色第一人称所见"),
    ("FPV", "First-Person View", "第一人称飞行", "穿越机视角，高速飞行"),
]

for i, (abbr, en, cn, note) in enumerate(shot_abbr):
    r = note_start + 2 + i
    ws.cell(row=r, column=1, value=abbr).font = Font(name="Microsoft YaHei", bold=True, size=10, color="C55A11")
    ws.cell(row=r, column=2, value=en).font = note_font
    ws.cell(row=r, column=3, value=cn).font = note_font
    ws.cell(row=r, column=4, value=note).font = note_font
    for col in range(1, 5):
        ws.cell(row=r, column=col).alignment = wrap_align
        ws.cell(row=r, column=col).border = thin_border

combo_start = note_start + 2 + len(shot_abbr) + 1
ws.cell(row=combo_start, column=1, value="💡 组合写法示例：").font = note_title_font
ws.merge_cells(start_row=combo_start, start_column=1, end_row=combo_start, end_column=5)

combo_examples = [
    ("close-up to wide shot", "从特写拉到远景（如：拉镜头、航拍拉升）"),
    ("medium shot to close-up", "从中景推到特写（如：推镜头、变焦推进）"),
    ("medium shot to extreme close-up", "从中景推到大特写（如：推至细节、急速冲撞变焦）"),
    ("low medium shot to eye-level close-up", "从低角度中景到平视近景（如：推+升复合运镜）"),
]

for i, (expr, desc) in enumerate(combo_examples):
    r = combo_start + 1 + i
    ws.cell(row=r, column=1, value=expr).font = Font(name="Microsoft YaHei", bold=True, size=10, color="C55A11")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=3, value=desc).font = note_font
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    for col in range(1, 6):
        ws.cell(row=r, column=col).alignment = wrap_align
        ws.cell(row=r, column=col).border = thin_border

# ============================================================
# Sheet 2: 运镜分类总表
# ============================================================
ws2 = wb.create_sheet("运镜分类总表")

taxonomy_headers = ["大类", "子类", "中文定义", "英文提示词", "适用场景", "速度建议"]
for col, h in enumerate(taxonomy_headers, 1):
    cell = ws2.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

taxonomy_widths = [14, 18, 40, 28, 24, 12]
for i, w in enumerate(taxonomy_widths, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

taxonomy_data = [
    ("固定镜头", "纯固定", "镜头完全不动，只有画面内元素或背景变动", "static shot, locked-off", "口播、对话、展示", "无"),
    ("固定镜头", "微固定(呼吸感)", "几乎不动，带极轻微呼吸感抖动", "static with subtle breathing", "真实感口播", "无"),
    ("推镜头", "缓慢推进", "镜头缓慢向主体靠近，聚焦注意力", "slow dolly in", "情绪递进、产品展示", "慢"),
    ("推镜头", "快速推近", "镜头快速向主体靠近，冲击感", "fast push in", "动作、惊吓、强调", "快"),
    ("推镜头", "微推", "几乎不可察觉的推进，增加张力", "subtle push in", "对话、口播", "极慢"),
    ("推镜头", "推至特写", "从中景推到局部特写", "dolly in to close-up", "产品、表情特写", "慢~中"),
    ("推镜头", "推至细节", "推到物体极近细节", "push in to detail", "产品纹理、文字", "慢"),
    ("推镜头", "阶梯推", "推-停-推，分段推进", "stepped dolly in", "节奏感叙事", "变速"),
    ("拉镜头", "缓慢拉远", "镜头缓慢远离主体", "slow dolly out", "揭示环境、结束", "慢"),
    ("拉镜头", "快速拉远", "镜头快速远离，震撼揭示", "fast pull out", "戏剧性揭示", "快"),
    ("拉镜头", "拉远后固定", "拉到一定距离后停住", "dolly out then hold", "转场前铺垫", "慢→停"),
    ("拉镜头", "拉至远景", "从特写一直拉到大远景", "pull out to wide shot", "环境揭示", "慢~中"),
    ("拉镜头", "无限拉远", "持续远离至极远", "infinite pull out", "航拍感、结局", "慢"),
    ("横移镜头", "向左平移", "匀速向左水平移动", "truck left", "场景扫描、揭示", "中"),
    ("横移镜头", "向右平移", "匀速向右水平移动", "truck right", "场景扫描、揭示", "中"),
    ("横移镜头", "斜向移动", "对角线方向移动", "diagonal tracking", "动态构图", "中"),
    ("横移镜头", "不规则移动", "非匀速非直线横移", "irregular lateral", "纪实、手持风格", "变速"),
    ("横移镜头", "弧线移动", "沿弧线横向移动", "arc tracking", "创意构图", "中"),
    ("横移镜头", "穿越移动", "穿过障碍物/人群横移", "tracking through", "探店、场景穿梭", "中~快"),
    ("横移镜头", "遮挡横移", "藏在柱子/门后，制造悬疑惊喜", "obstacle tracking", "悬疑、创意转场", "中"),
    ("环绕镜头", "顺时针环绕", "从左侧到右侧环绕", "clockwise orbit", "人物展示、产品", "慢~中"),
    ("环绕镜头", "逆时针环绕", "从右侧到左侧环绕", "counterclockwise orbit", "人物展示", "慢~中"),
    ("环绕镜头", "半环绕180°", "绕主体半圈", "180-degree orbit", "人物/产品展示", "中"),
    ("环绕镜头", "全环绕360°", "绕主体一整圈", "360-degree orbit", "全方位展示", "中"),
    ("环绕镜头", "缓慢弧形环绕", "缓慢弧线转换视角，唯美治愈", "slow arc orbit", "唯美展示、治愈", "慢"),
    ("环绕镜头", "螺旋上升环绕", "边环绕边升高", "spiral ascending orbit", "史诗感、揭示", "慢"),
    ("环绕镜头", "螺旋下降环绕", "边环绕边降低", "spiral descending orbit", "场景进入", "慢"),
    ("升降镜头", "升镜头", "从低位上升到高位，越升越广展现史诗感", "crane up, boom up", "揭示、情绪升华、史诗感", "慢~中"),
    ("升降镜头", "降镜头", "从高位下降到低位，聚焦人物增强故事感", "crane down, boom down", "场景代入、角色引入", "慢~中"),
    ("升降镜头", "升后俯瞰", "升到最高后俯视", "crane up to bird's eye", "全景揭示", "慢"),
    ("升降镜头", "降至平视", "从高处降到平视", "descend to eye level", "角色引入", "慢~中"),
    ("升降镜头", "升降+推拉", "升降同时推或拉", "crane + dolly", "复合电影感", "慢"),
    ("升降镜头", "揭示升降", "升起后揭示背后场景", "reveal crane shot", "悬念揭示", "慢~中"),
    ("摇镜头", "水平左摇", "镜头从右向左旋转", "pan left", "场景揭示、跟踪", "慢~中"),
    ("摇镜头", "水平右摇", "镜头从左向右旋转", "pan right", "场景揭示、跟踪", "慢~中"),
    ("摇镜头", "垂直上摇", "镜头从下向上旋转", "tilt up", "建筑、人物展示", "慢~中"),
    ("摇镜头", "垂直下摇", "镜头从上向下旋转", "tilt down", "场景进入、产品", "慢~中"),
    ("摇镜头", "仰拍上摇", "从脚拍到脸，提升气场", "low-angle tilt up", "穿搭展示、气场", "慢~中"),
    ("摇镜头", "俯拍下摇", "从上拍到全身，展示穿搭", "high-angle tilt down", "穿搭展示、全身", "慢~中"),
    ("摇镜头", "斜向摇", "对角线方向旋转", "diagonal pan-tilt", "动态构图", "中"),
    ("摇镜头", "360°旋转摇", "水平旋转一整圈", "360-degree pan", "全景环境展示", "中~慢"),
    ("摇镜头", "揭示摇", "摇到最终位置揭示关键元素", "reveal pan", "悬念、惊喜", "慢"),
    ("甩镜头", "水平甩", "快速水平甩向一侧", "horizontal whip pan", "转场、紧张", "极快"),
    ("甩镜头", "垂直甩", "快速垂直甩向上/下", "vertical whip tilt", "转场", "极快"),
    ("甩镜头", "甩接甩", "两个甩镜头衔接转场", "whip pan transition", "创意转场", "极快"),
    ("甩镜头", "对角甩", "斜向急速甩", "diagonal whip", "创意转场", "极快"),
    ("晃动镜头", "轻微晃动", "有呼吸感的微抖", "subtle handheld", "真实感、Vlog", "无"),
    ("晃动镜头", "明显晃动", "纪录片/战争场景晃动", "shaky handheld", "纪实、紧张", "无"),
    ("晃动镜头", "跟随晃动", "跟随主体同时晃动", "handheld tracking", "探店、Vlog", "中"),
    ("晃动镜头", "奔跑晃动", "模拟跑动拍摄", "running handheld", "追逐、紧张", "快"),
    ("晃动镜头", "手持纪实", "微抖有生活感纪录片感", "documentary handheld", "纪实、生活感", "无"),
    ("变焦镜头", "缓慢变焦推", "缓慢拉近焦距", "slow zoom in", "聚焦、强调", "慢"),
    ("变焦镜头", "快速变焦推", "瞬间拉近", "snap zoom in", "冲击、惊吓", "极快"),
    ("变焦镜头", "变焦拉", "焦距从长变短", "zoom out", "揭示、拉开距离", "慢~中"),
    ("变焦镜头", "希区柯克变焦", "推+反向变焦，眩晕感", "dolly zoom, vertigo", "恐惧、顿悟", "中"),
    ("变焦镜头", "微距变焦", "展现极致细节，科幻微观感", "macro zoom", "产品细节、微观", "慢"),
    ("变焦镜头", "急速冲撞变焦", "瞬间特写，表现震惊", "crash zoom", "冲击、惊吓", "极快"),
    ("变焦镜头", "平滑光学推进", "温柔拉近，突出面部情绪", "smooth optical zoom in", "情绪、人物", "慢"),
    ("变焦镜头", "平滑光学拉远", "慢慢拉宽，显环境空旷", "smooth optical zoom out", "环境揭示", "慢"),
    ("滚转镜头", "顺时针滚转", "画面顺时针旋转", "clockwise roll", "梦幻、迷幻", "慢"),
    ("滚转镜头", "逆时针滚转", "画面逆时针旋转", "counterclockwise roll", "梦幻、不安", "慢"),
    ("滚转镜头", "微滚转", "轻微倾斜增加不安感", "subtle dutch angle", "悬疑、不安", "极慢"),
    ("滚转镜头", "荷兰倾斜角", "画面歪斜，表现紧张不安", "dutch angle", "悬疑、紧张", "无"),
    ("跟随镜头", "背跟", "镜头在背后跟随前进", "following from behind", "行走、探索", "中"),
    ("跟随镜头", "正跟", "镜头在正前方后退拍摄", "backward facing subject", "面对面口播行走", "中"),
    ("跟随镜头", "侧跟", "镜头在侧面平行移动", "side tracking", "并排行走", "中"),
    ("跟随镜头", "低角度跟", "低机位跟随", "low-angle tracking", "脚步、宠物", "中"),
    ("跟随镜头", "高角度跟", "俯视角度跟随", "overhead tracking", "人群、路线", "中"),
    ("跟随镜头", "环绕跟", "边跟随边环绕", "orbiting tracking", "动作、展示", "中"),
    ("跟随镜头", "第一人称行走", "沉浸式第一人称视角", "first-person walking", "沉浸、代入感", "中"),
    ("航拍运镜", "俯冲", "从高空快速下降至主体", "drone dive", "震撼开场", "快"),
    ("航拍运镜", "拉升", "从低位快速上升至鸟瞰", "drone ascend", "揭示全貌", "中~快"),
    ("航拍运镜", "航拍环绕", "高空绕主体飞行", "aerial orbit", "建筑、景点", "慢~中"),
    ("航拍运镜", "航拍跟随", "高空跟随主体移动", "aerial tracking", "车辆、人群", "中"),
    ("航拍运镜", "航拍平移", "高空水平飞行", "aerial lateral", "地貌扫描", "中"),
    ("航拍运镜", "无人机高空飞跃", "上帝视角飞越，开阔治愈", "aerial flyover", "开场、环境", "中"),
    ("航拍运镜", "史诗无人机揭示", "从山后/建筑后露出，惊艳开场", "epic drone reveal", "震撼开场", "慢~中"),
    ("航拍运镜", "上帝俯拍", "90度正俯视，构图高级", "top-down, bird's eye", "构图、创意", "无"),
    ("航拍运镜", "FPV穿越机", "第一人称穿越飞行", "FPV drone shot", "穿越建筑、极限", "快"),
    ("特殊视角", "过肩镜头", "从肩后拍摄增强代入感故事感", "over-the-shoulder (OTS)", "对话、采访", "无"),
    ("特殊视角", "鱼眼/猫眼视角", "画面扭曲产生压迫悬疑感", "fisheye, peephole", "悬疑、创意", "无"),
    ("焦点控制", "虚焦转清晰", "从模糊渐变清晰，营造回忆心动感", "rack focus in", "回忆、心动", "慢"),
    ("焦点控制", "焦点切换", "前后景焦点转移，突出重点", "rack focus shift", "对比、重点转移", "中"),
]

for idx, (cat, sub, cn_def, en_kw, scene, speed) in enumerate(taxonomy_data, 1):
    row = idx + 1
    ws2.cell(row=row, column=1, value=cat).font = normal_font
    ws2.cell(row=row, column=2, value=sub).font = normal_font
    ws2.cell(row=row, column=3, value=cn_def).font = normal_font
    ws2.cell(row=row, column=4, value=en_kw).font = normal_font
    ws2.cell(row=row, column=5, value=scene).font = normal_font
    ws2.cell(row=row, column=6, value=speed).font = normal_font
    for col in range(1, 7):
        ws2.cell(row=row, column=col).alignment = wrap_align
        ws2.cell(row=row, column=col).border = thin_border

# ============================================================
# Sheet 3: 转场分类
# ============================================================
ws3 = wb.create_sheet("转场分类")

trans_headers = ["转场类型", "中文描述", "英文提示词", "适用场景"]
for col, h in enumerate(trans_headers, 1):
    cell = ws3.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

trans_widths = [16, 40, 28, 30]
for i, w in enumerate(trans_widths, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

transitions = [
    ("自然过渡", "平滑衔接，让两段内容有关联性", "smooth transition", "连续叙事、时间连续"),
    ("淡入", "从黑/白屏渐渐显现画面", "fade in", "开场、新章节开始"),
    ("淡出", "画面渐渐消失至黑/白屏", "fade out, fade to black", "结尾、时间流逝"),
    ("淡入/淡出(场景转换)", "大场景转换或时间转换时使用", "fade transition", "大场景转换、时间转换"),
    ("淡入/淡出(结尾)", "影片结尾处逐渐黑屏", "fade to black ending", "影片结尾"),
    ("交叉叠化", "前后帧中间带有交叉溶解的效果", "cross dissolve", "时间流逝、回忆、梦境"),
    ("前景遮挡无缝转场", "前景物体遮住镜头后切换到新场景", "foreground wipe transition", "创意无缝转场"),
    ("甩转场", "利用甩镜头的运动模糊衔接", "whip pan transition", "快节奏剪辑、场景跳跃"),
    ("匹配剪辑", "前后画面形状或动作匹配衔接", "match cut", "创意转场、隐喻"),
    ("跳切", "同景别直接跳跃剪辑", "jump cut", "口播、时间压缩、节奏感"),
    ("遮罩转场", "形状遮罩擦除切换", "mask wipe transition", "创意/MG风格"),
    ("闪白转场", "画面闪白后切换", "flash transition", "拍照、闪光、高能瞬间"),
    ("模糊转场", "画面虚焦后切换到新场景", "blur transition", "梦境、回忆、时间跳跃"),
    ("穿越镜头转场", "穿过缝隙/物体实现空间切换", "through-object transition", "创意无缝转场、空间穿越"),
    ("虚焦转场", "虚焦→清晰切换到新场景，营造回忆心动感", "focus transition", "回忆、梦境、场景切换"),
]

for idx, (name, desc, en, scene) in enumerate(transitions, 1):
    row = idx + 1
    ws3.cell(row=row, column=1, value=name).font = normal_font
    ws3.cell(row=row, column=2, value=desc).font = normal_font
    ws3.cell(row=row, column=3, value=en).font = normal_font
    ws3.cell(row=row, column=4, value=scene).font = normal_font
    for col in range(1, 5):
        ws3.cell(row=row, column=col).alignment = wrap_align
        ws3.cell(row=row, column=col).border = thin_border

# ============================================================
# Sheet 4: 口播场景运镜方案
# ============================================================
ws4 = wb.create_sheet("口播场景运镜方案")

kb_headers = ["序号", "场景名称", "推荐运镜", "推荐景别", "速度", "场景描述"]
for col, h in enumerate(kb_headers, 1):
    cell = ws4.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

kb_widths = [5, 14, 30, 14, 8, 40]
for i, w in enumerate(kb_widths, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

kb_data = [
    (1, "正面直拍", "固定镜头 / 微推", "中近景~近景", "慢/无", "人物正对镜头讲述，最基础的口播方式"),
    (2, "半身侧面", "固定 + 轻微弧线移动", "中景", "慢", "人物侧面半身入镜，增加画面层次"),
    (3, "走路口播", "正跟/侧跟 + 稳定器", "中景~中近景", "中", "人物边走边说，增加动态感和场景代入"),
    (4, "坐姿口播", "固定 / 缓推", "中景~近景", "慢/无", "人物坐在椅子/沙发上讲述，轻松氛围"),
    (5, "车内口播", "固定(吸盘) / 轻微晃动", "近景", "无", "车内空间拍摄，利用车辆自然晃动"),
    (6, "户外站立", "固定 / 缓慢环绕", "中景", "慢", "人物站在户外场景中讲述，环境展示"),
    (7, "探店口播", "跟随 + 手持晃动 + 推拉", "中景~全景", "中~快", "边走边逛边讲，推门进入，展示店内"),
    (8, "屏幕录制+画中画", "固定(人物) + 屏幕录制", "近景(人物)", "无", "人物小窗+屏幕大画面，教程类"),
    (9, "产品展示", "固定(人) + 推至产品特写", "中景→特写", "慢", "先拍人物说话，再推进到产品细节"),
    (10, "Vlog式", "手持自拍 + 晃动 + 场景切换", "近景~中近景", "中", "自拍视角，真实感，多场景切换"),
    (11, "采访式", "固定/缓推 + 正反打", "中景~近景", "慢", "类似采访构图，一问一答"),
    (12, "双人对话", "固定 + 摇镜头切换", "中景", "慢", "两人对话，摇镜头在两人间切换"),
    (13, "情景演绎", "跟随 + 推拉 + 升降组合", "多景别切换", "变速", "有剧情的情景短剧式口播"),
]

for idx, (num, name, yunjing, jingbie, speed, desc) in enumerate(kb_data, 1):
    row = idx + 1
    ws4.cell(row=row, column=1, value=num).font = normal_font
    ws4.cell(row=row, column=2, value=name).font = normal_font
    ws4.cell(row=row, column=3, value=yunjing).font = normal_font
    ws4.cell(row=row, column=4, value=jingbie).font = normal_font
    ws4.cell(row=row, column=5, value=speed).font = normal_font
    ws4.cell(row=row, column=6, value=desc).font = normal_font
    for col in range(1, 7):
        ws4.cell(row=row, column=col).alignment = wrap_align
        ws4.cell(row=row, column=col).border = thin_border

# ============================================================
# Sheet 5: 运镜识别度排名 (AI Platform Recognition Ranking)
# ============================================================
ws5 = wb.create_sheet("运镜识别度排名")

rank_headers = [
    "排名", "识别等级", "运镜关键词(中)", "Prompt关键词(英)",
    "Seedance", "Kling", "全网热度",
    "数据来源", "Prompt写法建议"
]
for col, h in enumerate(rank_headers, 1):
    cell = ws5.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

rank_widths = [5, 10, 16, 30, 9, 9, 10, 36, 48]
for i, w in enumerate(rank_widths, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

s_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
a_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
b_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
c_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
d_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

tier_fills = {"S": s_fill, "A": a_fill, "B": b_fill, "C": c_fill, "D": d_fill}

# (排名, 等级, 中文, 英文关键词, Seedance, Kling, 热度, 来源, 建议)
rank_data = [
    (1, "S", "推镜头(推/拉)", "dolly in / dolly out",
     "✅ 核心", "✅ 核心", "★★★★★",
     "Seedance官方9大运镜之一; 所有教程/帖子必提; Atlabs/ZSky/Medium等均列首位",
     "slow dolly in / smooth dolly forward — 被称为\"最万能的单一运镜\"，学第一个就学它"),
    (2, "S", "水平摇(左/右摇)", "pan left / pan right",
     "✅ 核心", "✅ 核心", "★★★★★",
     "Seedance官方9大运镜之一; 与dolly、zoom并称\"三大基础运镜覆盖80%需求\"",
     "slow pan right / smooth pan left — 用于揭示宽场景或跟随水平动作"),
    (3, "S", "变焦(推/拉)", "zoom in / zoom out",
     "✅ 核心", "✅ 核心", "★★★★★",
     "Seedance官方9大运镜之一; 三大基础之一; 但多篇指南警告\"最容易过度使用\"",
     "slow zoom in / subtle zoom out — 注意加速度修饰词，否则AI可能随机变焦"),
    (4, "S", "固定镜头", "static shot / locked-off",
     "✅ 核心", "✅ 核心", "★★★★☆",
     "所有指南列为基础; 不加运镜指令时AI默认输出static; Atlabs称其\"像监控摄像头\"",
     "static shot, tripod, no camera movement — 需明确指定否则AI默认给出的效果很平"),
    (5, "A", "垂直摇(上/下摇)", "tilt up / tilt down",
     "✅ 核心", "✅ 支持", "★★★★☆",
     "Seedance官方9大运镜之一; Atlabs列为10大必学第2位; 适合建筑/产品揭示",
     "smooth tilt up / slow tilt down — 配合 no pan, no dolly 可获得纯净的垂直摇"),
    (6, "A", "跟随镜头", "tracking shot",
     "✅ 核心", "✅ 核心", "★★★★☆",
     "Seedance官方9大运镜之一; TikTok/YouTube教程高频; Atlabs列第3位",
     "smooth tracking shot following subject / stabilized tracking — 适合行走/奔跑场景"),
    (7, "A", "升降镜头", "crane up / crane down",
     "✅ 核心", "✅ 支持", "★★★★☆",
     "Seedance官方9大运镜之一; 多篇指南列为\"史诗感必备\"; boom up等效",
     "crane up revealing wide landscape / smooth crane down to eye level"),
    (8, "A", "环绕镜头", "orbit / 360-degree orbit",
     "✅ 核心", "✅ 核心", "★★★★☆",
     "Seedance官方9大运镜之一; Atlabs列第8位; 适合人物/产品全方位展示",
     "full 360-degree orbit around subject / slow arc shot 180 degrees"),
    (9, "A", "晃动镜头(手持)", "handheld / handheld shake",
     "✅ 核心", "✅ 支持", "★★★★☆",
     "Seedance官方9大运镜之一; 所有指南均提及; 需配修饰词控制强度",
     "subtle handheld / frantic handheld with natural shake — 加 subtle 或 mild 控制力度"),
    (10, "A", "横移镜头", "truck left / truck right",
     "✅ 支持", "✅ 支持", "★★★☆☆",
     "Seedance完整词表列出; 多篇指南提及; 与tracking有区别(truck是平移不跟人)",
     "camera trucks left at steady pace / lateral tracking — 平行移动扫描场景"),
    (11, "A", "甩镜头", "whip pan",
     "✅ 核心", "✅ 支持", "★★★★☆",
     "Seedance官方9大运镜之一; Atlabs列第7位; X/TikTok热门转场技巧",
     "extremely fast horizontal whip pan with motion blur — 适合场景切换和音乐视频"),
    (12, "B", "推至特写", "push in / dolly in to close-up",
     "✅ 支持", "✅ 支持", "★★★☆☆",
     "Seedance词表列出push in; 多篇指南作为dolly变体提及",
     "slow push in to close-up / subtle push in on face — dolly in的常见具体用法"),
    (13, "B", "拉远揭示", "pull out / dolly out to wide",
     "✅ 支持", "✅ 支持", "★★★☆☆",
     "Seedance词表列出pull out; 作为dolly out变体常见",
     "pull out to reveal wide scene / smooth pull back — 用于揭示环境或结束镜头"),
    (14, "B", "弧线移动", "arc shot",
     "✅ 支持", "✅ 支持", "★★★☆☆",
     "Seedance词表列出; 被视为orbit的部分版本(90-180度)",
     "smooth arc shot 90 degrees around subject — 比完整orbit更常用，适合人物揭示"),
    (15, "B", "希区柯克变焦", "dolly zoom / vertigo effect",
     "✅ 支持", "⚠️ 部分", "★★★☆☆",
     "Seedance/Atlabs/Medium等多篇指南专门提及; TikTok上有大量教程",
     "dolly zoom on character's face, background warping — 需描述\"推镜+反向变焦\"效果"),
    (16, "B", "急速冲撞变焦", "crash zoom",
     "✅ 支持", "⚠️ 部分", "★★★☆☆",
     "Atlabs列第6位; 多篇指南提及; 用于震惊/冲击瞬间",
     "very fast optical zoom in with abrupt stop framing face — 限制在3秒内保持冲击力"),
    (17, "B", "焦点切换", "rack focus / rack focus shift",
     "✅ 支持", "⚠️ 部分", "★★★☆☆",
     "多篇AI视频指南列为\"特殊效果\"; YouTube教程常见",
     "rack focus from foreground to background subject — 需明确说明焦点从哪到哪"),
    (18, "B", "航拍/无人机", "aerial / drone shot",
     "⚠️ 部分", "✅ 支持", "★★★☆☆",
     "Seedance作为style modifier; 多篇指南提及aerial rising/flyover",
     "aerial rising shot over city / drone flyover — 作为风格修饰词而非独立运镜类型"),
    (19, "B", "第一人称视角", "POV / first-person",
     "⚠️ 部分", "⚠️ 部分", "★★★☆☆",
     "Seedance列为style modifier; Atlabs列第10位; TikTok沉浸式视频热门",
     "POV shot running through forest / first-person perspective walking — 需明确描述视角"),
    (20, "B", "荷兰倾斜角", "dutch angle",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "Seedance列为style modifier但不在9大核心运镜内; 悬疑/心理惊悚类常用",
     "dutch angle tracking shot, tilted frame — 作为风格修饰词使用效果更好"),
    (21, "C", "缓慢弧形环绕", "slow arc orbit",
     "✅ 支持", "⚠️ 部分", "★★☆☆☆",
     "属于orbit变体; 唯美类视频常用; 部分教程提及",
     "slow arc orbit around subject, dreamy — 用slow arc shot替代可能识别度更高"),
    (22, "C", "螺旋上升/下降", "spiral ascending/descending",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "属于orbit+crane组合; 少数教程提及; 作为复合运镜描述",
     "orbit around subject while gradually rising — 分解为两个动作描述识别度更高"),
    (23, "C", "仰拍上摇", "low-angle tilt up",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "tilt up变体; 穿搭展示博主常用; 需额外说明起始角度",
     "low-angle tilt up from feet to face — 需明确 low-angle 起始位置"),
    (24, "C", "俯拍下摇", "high-angle tilt down",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "tilt down变体; 需额外说明起始角度",
     "high-angle tilt down from face to full body — 同上需描述起始角度"),
    (25, "C", "遮挡横移", "obstacle tracking",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "创意运镜; 少数高级教程提及; 平台不一定识别专用术语",
     "camera tracks behind pillars, peek-a-boo reveal — 用自然语言描述动作效果更好"),
    (26, "C", "穿越移动", "tracking through",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "创意运镜; FPV类视频概念; 需详细描述穿越过程",
     "camera moves through narrow gap / tracking through crowd — 需描述穿越什么"),
    (27, "C", "微距变焦", "macro zoom",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "产品类视频使用; 需配合extreme close-up描述",
     "extreme close-up macro zoom on surface texture — 建议配合 macro lens 描述"),
    (28, "C", "平滑光学变焦", "smooth optical zoom in/out",
     "✅ 支持", "⚠️ 部分", "★★☆☆☆",
     "zoom in变体; 加smooth修饰词增强效果",
     "smooth optical zoom in on face — 本质是zoom in+smooth修饰词的组合"),
    (29, "C", "过肩镜头", "over-the-shoulder",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "对话场景常见; AI平台识别度不高; 需详细描述构图",
     "over-the-shoulder framing from behind person — 需明确描述两人位置关系"),
    (30, "C", "鱼眼视角", "fisheye / peephole",
     "⚠️ 部分", "❌ 弱", "★☆☆☆☆",
     "极少教程提及; 非标准AI运镜; 属于镜头效果而非运镜",
     "fisheye lens distortion, wide-angle — 作为镜头风格描述,不作为运镜指令"),
    (31, "D", "FPV穿越机", "FPV drone shot",
     "❌ 未列出", "⚠️ 部分", "★★☆☆☆",
     "不在Seedance官方9大运镜内; 概念热门但AI执行难度高",
     "high-speed first-person drone racing through obstacles — 可能需要分解为tracking+fast"),
    (32, "D", "滚转镜头", "roll / clockwise roll",
     "❌ 未列出", "❌ 弱", "★☆☆☆☆",
     "极少AI视频教程提及; 非标准AI运镜动作",
     "frame slowly rotates clockwise — 用自然语言描述旋转动作"),
    (33, "D", "上帝俯拍", "top-down / bird's eye",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "aerial变体; 需配合overhead/top-down描述",
     "90-degree overhead top-down perspective — 建议配合 aerial 使用"),
    (34, "D", "虚焦转清晰", "rack focus in",
     "⚠️ 部分", "⚠️ 部分", "★★☆☆☆",
     "rack focus变体; 需详细描述模糊→清晰过程",
     "transitioning from blurred bokeh to sharp focus on subject"),
    (35, "D", "阶梯推", "stepped dolly in",
     "❌ 弱", "❌ 弱", "★☆☆☆☆",
     "极小众; 几乎无AI教程提及此专用术语",
     "dolly forward, pause, then dolly forward again — 需拆解为多段动作描述"),
    (36, "D", "无限拉远", "infinite pull out",
     "❌ 弱", "❌ 弱", "★☆☆☆☆",
     "概念性描述; AI平台不识别此专用术语",
     "continuous slow dolly out to extreme wide shot — 用基础词+距离描述替代"),
    (37, "D", "揭示升降/揭示摇", "reveal crane / reveal pan",
     "⚠️ 部分", "❌ 弱", "★☆☆☆☆",
     "\"reveal\"作为修饰词可用但非独立运镜; 需描述揭示过程",
     "crane up to reveal landscape behind / pan right to reveal character"),
]

for idx, (rank, tier, cn, en, seed, kling, heat, source, tip) in enumerate(rank_data, 1):
    row = idx + 1
    ws5.cell(row=row, column=1, value=rank).font = normal_font
    tier_cell = ws5.cell(row=row, column=2, value=tier)
    tier_cell.font = Font(name="Microsoft YaHei", bold=True, size=12, color="000000")
    tier_cell.fill = tier_fills.get(tier, d_fill)
    tier_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=row, column=3, value=cn).font = normal_font
    ws5.cell(row=row, column=4, value=en).font = kw_font
    ws5.cell(row=row, column=5, value=seed).font = normal_font
    ws5.cell(row=row, column=6, value=kling).font = normal_font
    ws5.cell(row=row, column=7, value=heat).font = Font(name="Microsoft YaHei", size=10, color="C55A11")
    ws5.cell(row=row, column=8, value=source).font = Font(name="Microsoft YaHei", size=9)
    ws5.cell(row=row, column=9, value=tip).font = Font(name="Microsoft YaHei", size=9, color="2F5496")

    for col in range(1, 10):
        ws5.cell(row=row, column=col).alignment = wrap_align
        ws5.cell(row=row, column=col).border = thin_border

# Legend at the bottom
legend_start = len(rank_data) + 3
legends = [
    "📌 等级说明：",
    "S级 = 平台核心支持 + 全网高频出现，几乎所有AI视频教程/指南必提，Prompt识别率最高",
    "A级 = 平台明确支持，多数教程提及，Prompt识别率高，适合日常使用",
    "B级 = 平台可识别，部分教程提及，需配合修饰词使用效果更佳",
    "C级 = 平台部分识别，需要用自然语言描述辅助，识别效果因平台而异",
    "D级 = 平台识别度低或不支持，建议拆解为基础运镜组合描述",
    "",
    "📌 数据来源（2026年1-4月）：",
    "Seedance 2.0 官方Camera Movement Guide / Kling 3.0 官方Motion Control文档",
    "Atlabs AI《10 Camera Movements for Cinematic AI Videos》",
    "Seedance2.so《Camera Movement Prompts Complete Guide》(基于X @yyyole热帖 367赞575收藏)",
    "PromeAI《Seedance 2.0 Camera Movement Cheat Sheet》",
    "ZSky AI《50+ AI Video Prompts》/ Medium AI视频运镜专栏",
    "Androarena《38 Cinematic Camera Movements for Viral AI Videos》",
    "Reddit r/KlingAI_Videos / r/SoraAi 社区讨论",
]
for i, text in enumerate(legends):
    cell = ws5.cell(row=legend_start + i, column=1, value=text)
    cell.font = Font(name="Microsoft YaHei", size=9, color="666666") if i > 0 else note_title_font
    ws5.merge_cells(start_row=legend_start + i, start_column=1,
                    end_row=legend_start + i, end_column=9)

wb.save("/Users/edy/video-camera-movement/视频运镜分类与测试表.xlsx")
print("Excel saved successfully!")
